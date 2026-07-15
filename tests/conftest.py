import pytest

DIRTY_TEXT = "Contact Jane Smith at jane.smith@example.com or 555-123-4567. SSN: 123-45-6789."


@pytest.fixture
def dirty_pdf(tmp_path):
    import pypdf

    path = tmp_path / "dirty.pdf"
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.add_metadata({"/Title": "Dirty PDF", "/Author": "MetaMaker", "/Subject": DIRTY_TEXT})
    with open(path, "wb") as f:
        writer.write(f)
    return path


@pytest.fixture
def dirty_docx(tmp_path):
    import docx

    path = tmp_path / "dirty.docx"
    doc = docx.Document()
    doc.core_properties.author = "MetaMaker"
    doc.core_properties.title = "Dirty DOCX"
    doc.add_paragraph(DIRTY_TEXT)
    doc.save(path)
    return path


@pytest.fixture
def dirty_xlsx(tmp_path):
    import openpyxl

    path = tmp_path / "dirty.xlsx"
    wb = openpyxl.Workbook()
    wb.properties.creator = "MetaMaker"
    wb.properties.title = "Dirty XLSX"
    ws = wb.active
    ws.append(["name", "note"])
    ws.append(["Jane Smith", DIRTY_TEXT])
    wb.save(path)
    return path


@pytest.fixture
def dirty_jpg(tmp_path):
    import piexif
    from PIL import Image

    path = tmp_path / "dirty.jpg"
    img = Image.new("RGB", (32, 32), color="red")
    exif_dict = {"0th": {piexif.ImageIFD.Artist: b"MetaMaker"}}
    img.save(path, format="JPEG", exif=piexif.dump(exif_dict))
    return path


@pytest.fixture
def dirty_txt(tmp_path):
    path = tmp_path / "dirty.txt"
    path.write_text(DIRTY_TEXT, encoding="utf-8")
    return path


@pytest.fixture
def dirty_csv(tmp_path):
    path = tmp_path / "dirty.csv"
    path.write_text(f"name,note\nJane Smith,{DIRTY_TEXT}\n", encoding="utf-8")
    return path
